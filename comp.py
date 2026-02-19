import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, Side
from datetime import datetime
from openpyxl.utils import get_column_letter
import copy


# ---------------------------------------------------------
# APPLY STYLE FUNCTION
# ---------------------------------------------------------

def apply_style(path):

    wb = load_workbook(path)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Iterate over all worksheets
    for ws in wb.worksheets:

        # Set column width
        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 18

        # Apply border to all cells
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                min_col=1, max_col=ws.max_column):
            for cell in row:
                if cell.has_style:
                    cell.border = copy.copy(cell.border)

                cell.border = thin_border

    wb.save(path)


# ---------------------------------------------------------
# UPDATE EXECUTION REPORT FUNCTION
# ---------------------------------------------------------

def update_execution_report(path, BuildNumber, Trigger, Iteration):

    # Load data from Excel
    execution_report_data = pd.read_excel(path, sheet_name="Sheet1")
    status_data = pd.read_excel(path, sheet_name="Summary")

    # Get current date
    current_date = datetime.today().strftime("%Y-%m-%d")

    # Ensure current date column exists in execution sheet
    if current_date not in execution_report_data.columns:
        raise ValueError(f"{current_date} column not found in Sheet1")

    # Calculate benchmark data
    total_ts = len(execution_report_data[current_date].dropna())

    passed_ts = len(
        execution_report_data[
            execution_report_data[current_date].astype(str).str.startswith('Pass')
        ]
    )

    failed_ts = len(
        execution_report_data[
            execution_report_data[current_date].astype(str).str.startswith('Fail')
        ]
    )

    verification_ts = len(
        execution_report_data[
            execution_report_data[current_date].astype(str).str.startswith('VERI')
        ]
    )

    not_executed_ts = len(
        execution_report_data[
            execution_report_data[current_date].astype(str).str.startswith('Not')
        ]
    )

    # ---------------------------------------------------------
    # UPDATE SUMMARY SHEET STRUCTURE
    # ---------------------------------------------------------

    # Required fixed row labels
    required_rows = [
        "Date",
        "BuildNumber",
        "Trigger",
        "Iteration",
        "KITE Tool Version",
        "Total TS",
        "Pass",
        "Fail",
        "VERIFYMANUALLY",
        "Not Executed"
    ]

    # Ensure first column contains required rows
    if "Metric" not in status_data.columns:
        status_data = pd.DataFrame({"Metric": required_rows})
    else:
        status_data = status_data.set_index("Metric")
        for row in required_rows:
            if row not in status_data.index:
                status_data.loc[row] = ""
        status_data = status_data.reset_index()

    # Add new date column if not present
    if current_date not in status_data.columns:
        status_data[current_date] = ""

    # Fill values under current date column
    status_data.loc[status_data["Metric"] == "Date", current_date] = current_date
    status_data.loc[status_data["Metric"] == "BuildNumber", current_date] = BuildNumber
    status_data.loc[status_data["Metric"] == "Trigger", current_date] = Trigger
    status_data.loc[status_data["Metric"] == "Iteration", current_date] = Iteration
    status_data.loc[status_data["Metric"] == "Total TS", current_date] = total_ts
    status_data.loc[status_data["Metric"] == "Pass", current_date] = passed_ts
    status_data.loc[status_data["Metric"] == "Fail", current_date] = failed_ts
    status_data.loc[status_data["Metric"] == "VERIFYMANUALLY", current_date] = verification_ts
    status_data.loc[status_data["Metric"] == "Not Executed", current_date] = not_executed_ts

    # Write back to Excel
    with pd.ExcelWriter(path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        status_data.to_excel(writer, sheet_name="Summary", index=False)

    # Apply formatting
    apply_style(path)
